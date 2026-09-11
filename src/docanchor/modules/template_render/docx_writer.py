"""DOCX写入工具：基于python-docx写入干净DOCX。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import _Cell

from docanchor.common.schema import GlobalNode, NodeType, TableExtension
from docanchor.modules.template_render.style_map import StyleSpec, get_style


def _apply_run_style(run, style: StyleSpec) -> None:
    """应用样式到run。"""
    run.font.size = Pt(style.font_size_pt)
    if style.font_name:
        run.font.name = style.font_name
    if style.font_name_cn:
        # 设置东亚字体（中文）
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), style.font_name_cn)
        rfonts.set(qn("w:ascii"), style.font_name)
        rfonts.set(qn("w:hAnsi"), style.font_name)
    if style.bold:
        run.font.bold = True


def _apply_paragraph_style(paragraph, style: StyleSpec) -> None:
    """应用样式到段落。"""
    pf = paragraph.paragraph_format
    pf.line_spacing = style.line_spacing
    pf.space_before = Pt(style.space_before_pt)
    pf.space_after = Pt(style.space_after_pt)
    if style.first_line_indent_chars > 0:
        pf.first_line_indent = Pt(style.font_size_pt * style.first_line_indent_chars)
    if style.keep_with_next:
        pf.keep_with_next = True
    if style.page_break_before:
        pf.page_break_before = True
    if style.alignment == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif style.alignment == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif style.alignment == "both":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_styled_paragraph(
    doc,
    text: str,
    style: StyleSpec,
) -> None:
    """添加应用样式的段落。"""
    p = doc.add_paragraph()
    _apply_paragraph_style(p, style)
    run = p.add_run(text)
    _apply_run_style(run, style)


def _add_styled_table(
    doc,
    table_ext: TableExtension,
    style: StyleSpec,
) -> None:
    """添加应用样式的表格。"""
    rows = table_ext.row_count
    cols = table_ext.col_count
    if rows <= 0 or cols <= 0:
        return

    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 填充单元格
    matrix = table_ext.cell_matrix or []
    for r_idx in range(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            text = matrix[r_idx][c_idx] if r_idx < len(matrix) and c_idx < len(matrix[r_idx]) else ""
            cell.text = text
            # 应用样式
            for paragraph in cell.paragraphs:
                _apply_paragraph_style(paragraph, style)
                for run in paragraph.runs:
                    _apply_run_style(run, style)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 重复表头（第一行）
    if table_ext.repeat_header and rows > 0:
        first_row = table.rows[0]
        # 设置tr的tblHeader属性
        from docx.oxml import OxmlElement
        trPr = first_row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)


def render_node(
    doc,
    node: GlobalNode,
) -> None:
    """渲染单个节点到docx文档。

    Args:
        doc: python-docx Document。
        node: 待渲染节点。
    """
    if node.node_type == NodeType.DOCUMENT_ROOT:
        return

    if node.node_type == NodeType.HEADING:
        style = get_style(NodeType.HEADING, node.global_level)
        # 优先使用内置Heading样式（python-docx按name匹配）
        heading_name = f"Heading {node.global_level}"
        try:
            p = doc.add_paragraph()
            _apply_paragraph_style(p, style)
            # 设置内置样式
            if heading_name in [s.name for s in doc.styles]:
                p.style = doc.styles[heading_name]
            run = p.add_run(node.text)
            _apply_run_style(run, style)
            return
        except (KeyError, ValueError):
            pass
        _add_styled_paragraph(doc, node.text, style)

    elif node.node_type == NodeType.PARAGRAPH:
        style = get_style(NodeType.PARAGRAPH)
        _add_styled_paragraph(doc, node.text, style)

    elif node.node_type == NodeType.LIST:
        style = get_style(NodeType.LIST)
        _add_styled_paragraph(doc, node.text, style)

    elif node.node_type == NodeType.TABLE:
        style = get_style(NodeType.TABLE)
        if node.table:
            _add_styled_table(doc, node.table, style)

    elif node.node_type == NodeType.IMAGE:
        style = get_style(NodeType.IMAGE)
        # 图片：占位文字
        p = doc.add_paragraph()
        _apply_paragraph_style(p, style)
        run = p.add_run(f"[图片: {node.text or '(空)'}]")
        _apply_run_style(run, style)

    elif node.node_type == NodeType.CAPTION:
        style = get_style(NodeType.CAPTION)
        _add_styled_paragraph(doc, node.text, style)

    elif node.node_type == NodeType.FORMULA:
        style = get_style(NodeType.FORMULA)
        _add_styled_paragraph(doc, node.text, style)


__all__ = ["render_node"]