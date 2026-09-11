"""模板灌注与文档重建单测。"""

from __future__ import annotations

from pathlib import Path

from docanchor.common.idgen import new_node_id, reset_counters
from docanchor.common.schema import (
    DocumentTree,
    GlobalNode,
    NodeType,
    TableExtension,
)
from docanchor.modules.template_render.builder import render_document_tree


def _simple_tree() -> DocumentTree:
    """构造一个简单文档树用于测试。"""
    reset_counters()
    root = GlobalNode(node_id="root", node_type=NodeType.DOCUMENT_ROOT)
    h1 = GlobalNode(
        node_id="h1",
        node_type=NodeType.HEADING,
        text="第一章 测试",
        global_level=1,
        page_range=[1, 1],
    )
    p1 = GlobalNode(
        node_id="p1",
        node_type=NodeType.PARAGRAPH,
        text="正文段落",
        page_range=[1, 1],
    )
    tbl = GlobalNode(
        node_id="tbl1",
        node_type=NodeType.TABLE,
        text="",
        page_range=[1, 1],
        table=TableExtension(
            row_count=2,
            col_count=2,
            cell_matrix=[["A", "B"], ["1", "2"]],
        ),
    )
    root.children_ids = [h1.node_id]
    h1.children_ids = [p1.node_id, tbl.node_id]
    p1.parent_id = h1.node_id
    tbl.parent_id = h1.node_id
    return DocumentTree(
        document_id="t1",
        root=root,
        nodes={
            "root": root,
            "h1": h1,
            "p1": p1,
            "tbl1": tbl,
        },
    )


class TestRenderDocumentTree:
    def test_render_creates_docx(self, tmp_path: Path) -> None:
        tree = _simple_tree()
        out_path = tmp_path / "output.docx"
        result_path = render_document_tree(tree, output_path=out_path)
        assert result_path == out_path
        assert out_path.exists()

    def test_render_heading_has_style(self, tmp_path: Path) -> None:
        from docx import Document

        tree = _simple_tree()
        out_path = tmp_path / "output.docx"
        render_document_tree(tree, output_path=out_path)
        doc = Document(str(out_path))
        # 第一个段落应是Heading 1
        h1_paragraphs = [
            p for p in doc.paragraphs if p.style and "Heading 1" in p.style.name
        ]
        assert len(h1_paragraphs) >= 1
        assert "第一章 测试" in h1_paragraphs[0].text

    def test_render_table_has_cells(self, tmp_path: Path) -> None:
        from docx import Document

        tree = _simple_tree()
        out_path = tmp_path / "output.docx"
        render_document_tree(tree, output_path=out_path)
        doc = Document(str(out_path))
        assert len(doc.tables) >= 1
        table = doc.tables[0]
        assert len(table.rows) == 2
        assert len(table.columns) == 2
        assert table.cell(0, 0).text == "A"
        assert table.cell(1, 1).text == "2"

    def test_render_creates_template_if_missing(self, tmp_path: Path) -> None:
        """模板不存在时自动创建默认模板。"""
        tree = _simple_tree()
        template_path = tmp_path / "templates" / "default.docx"
        assert not template_path.exists()
        out_path = tmp_path / "output.docx"
        render_document_tree(tree, template_path=template_path, output_path=out_path)
        assert template_path.exists()